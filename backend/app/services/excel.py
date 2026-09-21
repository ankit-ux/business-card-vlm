import io

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("first_name", "First Name"),
    ("last_name", "Last Name"),
    ("position", "Position / Job Title"),
    ("company", "Company"),
    ("location", "Location"),
    ("phone", "Phone Number"),
    ("email", "Email Address"),
    ("status", "Status"),
    ("is_duplicate", "Duplicate?"),
    ("filename", "Source File"),
]


def leads_to_excel(leads: list[dict]) -> bytes:
    rows = []
    for lead in leads:
        row = {label: lead.get(key, "") for key, label in COLUMNS}
        row["Duplicate?"] = "Yes" if lead.get("is_duplicate") else "No"
        rows.append(row)

    df = pd.DataFrame(rows, columns=[label for _, label in COLUMNS])

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Leads")
        ws = writer.sheets["Leads"]

        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx, _ in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].tolist()])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

        ws.freeze_panes = "A2"

    buffer.seek(0)
    return buffer.read()
